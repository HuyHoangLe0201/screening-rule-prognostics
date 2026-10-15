# -*- coding: utf-8 -*-
r"""The pre-registered external test of PREREG_p3_fatigue.md: the screen on
strain-controlled fatigue specimens, an asset class that took no part in
building it, with a decline predicted in advance.

    python code/study_p3_fatigue.py cache      read the workbooks once into .npz
    python code/study_p3_fatigue.py run        screen every cohort, predict every specimen
    python code/study_p3_fatigue.py summary    W1-W4

Every constant is the protocol's; the script refuses to run on any other
version of it.  The workbooks are read from $P3_DATA/fatigue_ti64/.

Two deviations, both fixed from the column headers, the Axial Count column
and the Table of Contents before any stress was read, and both reported:

* The workbooks list one reversal per row (Time, Axial Count in segments,
  Stress, Strain), not one cycle per row as the protocol assumed from the
  published description.  The count is rounded to an integer reversal number
  (one sheet counts in half cycles, with a median step of 0.5, and is
  doubled).  Cycle c holds reversals 2c - 1 and 2c; a cycle with exactly
  these two is one row of the protocol, with s_max, s_min, e_max and e_min the
  larger and smaller of its two stresses and strains.  Cycles with more rows,
  where the counter restarted, are dropped.
* Long records are cut at the workbook's row limit, so the last recorded
  cycle cannot identify a run-out.  The Table of Contents gives each
  specimen's reversals to failure, 2Nf, and a specimen is a run-out when
  2Nf / 2 is 10^6 or more, the article's definition.  The truncated records
  are exactly the run-outs.  One other record, sa 0.006 (3)R-0.5, stops 8 %
  short of its 2Nf / 2, so its last recorded cycle is not its life; a record
  ending more than 1 % short is excluded, and the verdict is also reported
  with it included as the protocol's wording would have it.
"""
from __future__ import annotations

import glob
import hashlib
import json
import math
import os
import re
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

PREREG = os.path.join(HERE, "PREREG_p3_fatigue.md")
FROZEN_SHA256 = "c7128b1cb4d2a033371a93fc4cc2e23cb04ce59c49643134c4d52d52b145a6ed"
FROZEN_AT = "2026-10-15T15:47:22+07:00"
DATA = os.environ.get("P3_DATA", os.path.join(ROOT, "data"))
SRC = os.path.join(DATA, "fatigue_ti64")
CACHE = os.path.join(SRC, "cache")
BE = os.path.join(ROOT, "results", "bearings")
ROWS = os.path.join(BE, "p3_fatigue_rows.json")
OUT = os.path.join(BE, "p3_fatigue.json")
POINTS = (0.50, 0.70, 0.85)
SEED, B_BOOT, GB_SEED = 20261023, 10_000, 20261020
RUNOUT, MIN_ROWS, MIN_COHORT = 10**6, 60, 4


def check_frozen():
    h = hashlib.sha256(open(PREREG, "rb").read()).hexdigest()
    if h != FROZEN_SHA256:
        sys.exit(f"the protocol has changed since it was frozen: {h}")
    return h


def stable_seed(*key):
    digest = hashlib.blake2b(repr(key).encode("utf-8"), digest_size=8).digest()
    return int.from_bytes(digest, "big") % 10000


def _ratio(fname):
    m = re.search(r"R=(-?[0-9]+(?:\.[0-9]+)?)", fname)
    return float(m.group(1))


def cache():
    """Read every specimen sheet once (Axial Count, Stress, Strain) into .npz.

    The column positions come from the header row ('Time', 'Axial Count',
    'Stress', 'Strain'); the Table of Contents gives each specimen's 2Nf.
    """
    import openpyxl
    check_frozen()
    os.makedirs(CACHE, exist_ok=True)
    toc_all, mapping = {}, {}
    for p in sorted(glob.glob(os.path.join(SRC, "*.xlsx"))):
        R = _ratio(os.path.basename(p))
        wb = openpyxl.load_workbook(p, read_only=True)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            if row and isinstance(row[0], str) and row[0].startswith("sa") and row[1] is not None:
                toc_all[row[0].strip()] = {"R": R, "2Nf": row[1]}
        for ws in wb.worksheets[1:]:
            rows = ws.iter_rows(values_only=True)
            head = None
            for r in rows:
                if r and "Axial Count" in r:
                    head = list(r)
                    break
            ic, i_s, i_e = head.index("Axial Count"), head.index("Stress"), head.index("Strain")
            cnt, s, e = [], [], []
            for r in rows:
                if r is None or len(r) <= max(ic, i_s, i_e):
                    continue
                c, ss, ee = r[ic], r[i_s], r[i_e]
                if all(isinstance(v, (int, float)) for v in (c, ss, ee)):
                    cnt.append(c), s.append(ss), e.append(ee)
            sid = ws.title.strip()
            np.savez_compressed(os.path.join(CACHE, re.sub(r"[^A-Za-z0-9.()+-]", "_", sid) + ".npz"),
                                count=np.array(cnt, float), stress=np.array(s, float),
                                strain=np.array(e, float))
            mapping[sid] = {"file": os.path.basename(p), "R": R,
                            "columns": {"count": ic, "stress": i_s, "strain": i_e}}
            print(f"  {sid}: {len(cnt)} reversals", flush=True)
        wb.close()
    json.dump({"toc": toc_all, "sheets": mapping}, open(os.path.join(CACHE, "index.json"), "w"),
              indent=1)


def load_specimens(keep_short=False):
    """Specimens keyed by sheet name, with the protocol's indicator and exclusions."""
    idx = json.load(open(os.path.join(CACHE, "index.json")))
    specs, excluded = {}, {}
    for sid, meta in sorted(idx["sheets"].items()):
        nf2 = float(idx["toc"][sid]["2Nf"])
        if nf2 / 2.0 >= RUNOUT:                               # the article's run-out
            excluded[sid] = "run-out"
            continue
        z = np.load(os.path.join(CACHE, re.sub(r"[^A-Za-z0-9.()+-]", "_", sid) + ".npz"))
        cnt, s, e = z["count"], z["stress"], z["strain"]
        if abs(np.median(np.diff(cnt)) - 0.5) < 0.1:           # counted in half cycles
            cnt = 2.0 * cnt - 1.0
        cnt = np.rint(cnt)
        cyc = np.ceil(cnt / 2.0).astype(np.int64)            # reversals 2c-1 and 2c
        keep = cnt >= 1
        cyc, s, e = cyc[keep], s[keep], e[keep]
        order = np.argsort(cyc, kind="stable")
        cyc, s, e = cyc[order], s[order], e[order]
        u, start, k = np.unique(cyc, return_index=True, return_counts=True)
        two = k == 2
        a, b = start[two], start[two] + 1
        t = u[two].astype(float)
        smax, smin = np.maximum(s[a], s[b]), np.minimum(s[a], s[b])
        emax, emin = np.maximum(e[a], e[b]), np.minimum(e[a], e[b])
        ok = (np.isfinite(smax) & np.isfinite(smin) & np.isfinite(emax) & np.isfinite(emin)
              & (smax - smin > 0) & (emax - emin > 0))
        t, smax, smin, emax, emin = t[ok], smax[ok], smin[ok], emax[ok], emin[ok]
        n = len(t)
        dropped = int(len(u) - n)
        if n == 0:
            excluded[sid] = "no complete cycle"
            continue
        life = float(t[-1])                                   # before any averaging
        if life >= RUNOUT:
            excluded[sid] = "run-out"
            continue
        if life < 0.99 * nf2 / 2.0 and not keep_short:
            excluded[sid] = "record ends before failure"
            continue
        if n < MIN_ROWS:
            excluded[sid] = "fewer than 60 rows"
            continue
        if n > 5000:                                          # the protocol's averaging
            bsz = math.ceil(n / 1000)
            m = n // bsz

            def blk(x):
                return x[:m * bsz].reshape(m, bsz).mean(axis=1)
            t, smax, smin, emax, emin = (blk(x) for x in (t, smax, smin, emax, emin))
            n = m
        E = (smax - smin) / (emax - emin)
        nh = min(40, max(5, n // 10))
        h = 1.0 - E / np.median(E[:nh])
        specs[sid] = {"cohort": meta["R"], "t": t, "h": h, "rul": life - t, "n": n,
                      "life": life, "dt": float(np.median(np.diff(t))),
                      "amp": float(np.median((emax - emin) / 2.0)),
                      "cycles_dropped": dropped, "two_Nf": nf2}
    return specs, excluded, idx


def _screen(fracs, rng):
    import study_p3_review as R
    if len(fracs) < MIN_COHORT:
        return {"n": len(fracs), "verdict": "too few"}
    x = np.array(fracs, float)
    x = np.where(x >= 1.0, np.inf, x)
    rate = float(np.isfinite(x).mean())
    idx = rng.integers(0, len(x), size=(B_BOOT, len(x)))
    lo, hi = np.quantile(np.median(x[idx], axis=1), [0.025, 0.975], method="inverted_cdf")

    def pct(z):
        return None if not np.isfinite(z) else 100 * float(z)
    verdict = "check indicator" if rate < 0.25 else R._verdict3(pct(lo), pct(hi), 70.0)
    return {"n": len(x), "rate": rate, "censored_median_pct": pct(np.median(x)),
            "ci": [pct(lo), pct(hi)], "verdict": verdict}


def _distinct(amps):
    a = sorted(amps)
    return 1 + sum(1 for x, y in zip(a, a[1:]) if y > 1.01 * x) if a else 0


def _baseline(specs, ref, k, t):
    """The protocol's lifetime baseline: a life table on lives rescaled to k's amplitude."""
    import study_p3_review as R
    lives = [specs[j]["life"] for j in ref]
    amps = [specs[j]["amp"] for j in ref]
    if _distinct(amps) < 3:
        return R._life_table(lives, t), False
    b, a = np.polyfit(np.log(amps), np.log(lives), 1)

    def lhat(x):
        return math.exp(a + b * math.log(x))
    scaled = [L * lhat(specs[k]["amp"]) / lhat(x) for L, x in zip(lives, amps)]
    return R._life_table(scaled, t), True


def _gbr(specs, keys_by_cohort):
    """The step protocol's comparator: features of h at index j, log RUL target."""
    from sklearn.ensemble import GradientBoostingRegressor
    from study_p3_step import _features
    out = {}
    for c, keys in keys_by_cohort.items():
        keys = sorted(keys)
        fold = {k: int(i) % 5 for k, i in zip(keys, np.random.default_rng(GB_SEED).permutation(len(keys)))}
        for f in range(5):
            test = [k for k in keys if fold[k] == f]
            if not test:
                continue
            X, y = [], []
            for k in keys:
                if fold[k] == f:
                    continue
                r = specs[k]
                for j in range(20, r["n"] - 1, 5):
                    if r["rul"][j] > 0:
                        X.append(_features(r["h"], j))
                        y.append(np.log(r["rul"][j]))
            if not X:
                continue
            models = {a: GradientBoostingRegressor(loss="quantile", alpha=a, n_estimators=200,
                                                   max_depth=3, learning_rate=0.05,
                                                   random_state=0).fit(X, y)
                      for a in (0.5, 0.025, 0.975)}
            for k in test:
                r = specs[k]
                for p in POINTS:
                    j = int(p * r["n"])
                    if j < 20 or j >= r["n"] - 1:
                        continue
                    x = [_features(r["h"], j)]
                    out[(k, j)] = tuple(float(np.exp(models[a].predict(x)[0])) for a in (0.5, 0.025, 0.975))
    return out


def run():
    check_frozen()
    import study_bearings as sb
    import study_p3_prereg as P
    specs, excluded, idx = load_specimens()
    for k, r in specs.items():
        i = P.NS["onset"](r["h"])
        r["onset_idx"] = int(i) if i > 0 else r["n"]
        r["onset_frac"] = i / r["n"] if i > 0 else 1.0
    cohorts = {}
    for k, r in specs.items():
        cohorts.setdefault(r["cohort"], []).append(k)
    rng = np.random.default_rng(SEED)
    screen = {str(c): _screen([specs[k]["onset_frac"] for k in sorted(ks)], rng)
              for c, ks in sorted(cohorts.items())}
    screen["pooled"] = _screen([specs[k]["onset_frac"] for k in sorted(specs)], rng)
    for c, s in screen.items():
        print(f"  R={c}: {s['n']} specimens -> {s['verdict']} "
              f"{s.get('censored_median_pct')} {s.get('ci')}", flush=True)
    # the short record kept, as the protocol's wording would have it
    s2 = load_specimens(keep_short=True)[0]
    short = sorted(set(s2) - set(specs))
    incl = {}
    for c in sorted({s2[k]["cohort"] for k in short}):
        ks = sorted(k for k in s2 if s2[k]["cohort"] == c)
        fr = []
        for k in ks:
            i = P.NS["onset"](s2[k]["h"])
            fr.append(i / s2[k]["n"] if i > 0 else 1.0)
        incl[str(c)] = _screen(fr, np.random.default_rng(SEED))
        print(f"  with {short}: R={c} -> {incl[str(c)]}", flush=True)
    screened = {c: ks for c, ks in cohorts.items() if screen[str(c)]["verdict"] != "too few"}
    gb = _gbr(specs, screened)
    rows = []
    for c, ks in sorted(screened.items()):
        for k in sorted(ks):
            r = specs[k]
            ref = [j for j in ks if j != k]
            D = float(np.median([np.median(specs[j]["h"][-3:]) for j in ref]))
            med = float(np.median([max(specs[j]["n"] - specs[j]["onset_idx"], 20) for j in ref]))
            rec = {"t": r["t"], "hi": r["h"], "condition": 0, "life": r["life"],
                   "n": r["n"], "dt": r["dt"]}
            for p in POINTS:
                j = int(p * r["n"])
                if j < 20 or j >= r["n"] - 1:
                    continue
                true = float(r["rul"][j])
                if true <= 0:
                    continue
                pr = sb.predict_one(rec, j, D, "BS", informed=True, coupling=True,
                                    med_life_steps=med, seed=stable_seed(k, j), n_part=600)
                (lt, scaled) = _baseline(specs, ref, k, float(r["t"][j]))
                g = gb.get((k, j))
                row = {"specimen": k, "cohort": c, "n": r["n"], "k": j, "at_life": p,
                       "onset_frac": r["onset_frac"], "true": true, "lt_scaled": scaled,
                       "lt_pred": lt[0], "lt_lo": lt[1], "lt_hi": lt[2]}
                if g:
                    row.update({"gb_pred": g[0], "gb_lo": g[1], "gb_hi": g[2]})
                if pr is not None:
                    row.update({"pred": pr["pred"], "lo": pr["lo"], "hi": pr["hi"]})
                rows.append(row)
            print(f"  R={c} {k}: onset {r['onset_frac']:.2f}, n {r['n']}", flush=True)
    json.dump({"screen": screen, "screen_with_short": incl, "rows": rows, "excluded": excluded,
               "sheets": idx["sheets"],
               "specimens": {k: {kk: r[kk] for kk in ("cohort", "n", "onset_idx", "onset_frac",
                                                      "life", "dt", "amp", "cycles_dropped",
                                                      "two_Nf")}
                             for k, r in specs.items()}},
              open(ROWS, "w"), indent=1)


def _ape(r, k="pred"):
    return 100.0 * abs(r[k] - r["true"]) / r["true"]


def _cov(r, lo="lo", hi="hi"):
    return float(r[lo] <= r["true"] <= r[hi])


def summary():
    h = check_frozen()
    S = json.load(open(ROWS))
    screen = S["screen"]
    rows = [r for r in S["rows"] if "pred" in r]
    rng = np.random.default_rng(SEED)
    by = {}
    for r in rows:
        by.setdefault(r["specimen"], []).append(r)

    def flat(us):
        return [r for u in us for r in by[u]]

    def score(rs):
        out = {"specimens": len({r["specimen"] for r in rs}), "predictions": len(rs),
               "medape_path": float(np.median([_ape(r) for r in rs])),
               "coverage_path": float(np.mean([_cov(r) for r in rs])),
               "medape_lt": float(np.median([_ape(r, "lt_pred") for r in rs])),
               "coverage_lt": float(np.mean([_cov(r, "lt_lo", "lt_hi") for r in rs]))}
        g = [r for r in rs if "gb_pred" in r]
        if g:
            out["medape_gb"] = float(np.median([_ape(r, "gb_pred") for r in g]))
            out["coverage_gb"] = float(np.mean([_cov(r, "gb_lo", "gb_hi") for r in g]))
        return out
    cohorts = sorted(c for c in {float(c) for c in screen if c != "pooled"})
    scr = [c for c in cohorts if screen[str(c)]["verdict"] != "too few"]
    res = {"protocol_sha256": h, "protocol_frozen_at": FROZEN_AT, "screen": screen,
           "screen_with_short": S["screen_with_short"], "excluded": S["excluded"],
           "specimens": S["specimens"],
           "by_cohort": {str(c): score([r for r in rows if r["cohort"] == c]) for c in scr
                         if any(r["cohort"] == c for r in rows)},
           "pooled": score(rows) if rows else None}
    dec = [c for c in scr if screen[str(c)]["verdict"] == "declines"]
    nd = [c for c in scr if screen[str(c)]["verdict"] == "not declined"]
    res["W1"] = {"holds": bool(dec) and not nd, "declined": dec, "not_declined": nd}
    res["W2"] = {"holds": all(screen[str(c)]["verdict"] != "check indicator" for c in scr),
                 "rates": {str(c): screen[str(c)].get("rate") for c in scr}}
    if dec:
        Ud = sorted(u for u in by if by[u][0]["cohort"] in dec)

        def st(us):
            fa = flat(us)
            return np.median([_ape(r) for r in fa]) - np.median([_ape(r, "lt_pred") for r in fa])
        obs = st(Ud)
        bt = [st([Ud[i] for i in rng.integers(0, len(Ud), len(Ud))]) for _ in range(B_BOOT)]
        c3 = np.percentile(bt, [2.5, 97.5]).tolist()
        res["W3"] = {"testable": True, "diff": float(obs), "ci": c3, "holds": bool(not c3[1] < 0)}
    else:
        res["W3"] = {"testable": False}
    late = sorted(u for u in by if by[u][0]["onset_frac"] >= 0.70)
    early = sorted(u for u in by if by[u][0]["onset_frac"] < 0.70)
    if len(late) >= 5 and len(early) >= 5:
        def cv(a, b):
            return np.mean([_cov(r) for r in flat(a)]) - np.mean([_cov(r) for r in flat(b)])
        obs = cv(late, early)
        bt = [cv([late[i] for i in rng.integers(0, len(late), len(late))],
                 [early[i] for i in rng.integers(0, len(early), len(early))]) for _ in range(B_BOOT)]
        c4 = np.percentile(bt, [2.5, 97.5]).tolist()
        res["W4"] = {"testable": True, "n_late": len(late), "n_early": len(early),
                     "diff": float(obs), "ci": c4, "holds": bool(c4[1] < 0)}
    else:
        res["W4"] = {"testable": False, "n_late": len(late), "n_early": len(early)}
    # after the fact, outside the protocol: onset against strain amplitude, and
    # how often the indicator had already crossed D at the prediction point
    sp = S["specimens"]
    bands = {"le_0.006": lambda a: a <= 0.0061, "0.007_to_0.010": lambda a: 0.0069 < a <= 0.0101,
             "0.012": lambda a: a > 0.011}
    res["exploratory"] = {
        "late_by_amplitude": {b: {"late": sum(1 for v in sp.values() if f(v["amp"]) and v["onset_frac"] >= 0.70),
                                  "of": sum(1 for v in sp.values() if f(v["amp"]))}
                              for b, f in bands.items()},
        "path_predicts_zero": sum(1 for r in rows if r["pred"] <= 0.0), "predictions": len(rows)}
    json.dump(res, open(OUT, "w", encoding="utf-8"), indent=1)
    print("exploratory", res["exploratory"])
    for c, v in res["by_cohort"].items():
        s = screen[c]
        print(f"R={c}: onset {s['censored_median_pct']} {s['ci']} {s['verdict']:13s} "
              f"path {v['medape_path']:.1f}/{v['coverage_path']:.2f}  LT {v['medape_lt']:.1f}/"
              f"{v['coverage_lt']:.2f}  GB {v.get('medape_gb', float('nan')):.1f}/"
              f"{v.get('coverage_gb', float('nan')):.2f}")
    print("pooled screen", screen["pooled"])
    print("excluded", S["excluded"])
    for x in ("W1", "W2", "W3", "W4"):
        print(x, res[x])


if __name__ == "__main__":
    {"cache": cache, "run": run, "summary": summary}.get(sys.argv[1] if len(sys.argv) > 1 else "",
                                                         lambda: sys.exit(__doc__))()
